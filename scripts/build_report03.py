"""Ghép lý thuyết từ hai PDF, chỉ bổ sung/sửa nội dung ở màu FF0000.

Tạo Word với heading/TOC field và OMML; PDF được xuất từ chính bản Word.
Yêu cầu python-docx, PyMuPDF và pandoc (để chuyển LaTeX sang OMML).
"""
from pathlib import Path
from collections import Counter
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import csv
import json
import re
import subprocess
import unicodedata
import math
from difflib import SequenceMatcher
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from report_details03 import add_intro,add_details

ROOT=Path(__file__).resolve().parents[1]
ART=ROOT/'artifacts_tl03';DATA=ART/'data';FIG=ART/'figures'
NAME='tieuluan03_06_02_giangkt.254'
BLACK='000000';RED='FF0000'
S1=json.loads((ART/'source/tl01.json').read_text())
S2=json.loads((ART/'source/tl02.json').read_text())
TAB=json.loads((DATA/'tabular_results.json').read_text())
CNN=json.loads((DATA/'cnn_results.json').read_text())
SEQ=json.loads((DATA/'sequence_results.json').read_text())
doc=Document();section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.left_margin=Cm(3);section.right_margin=Cm(2)
section.top_margin=Cm(2.5);section.bottom_margin=Cm(2.5)
section.header_distance=Cm(1.1);section.footer_distance=Cm(1.2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2','TOC 3']:
    try:s=doc.styles[name]
    except KeyError:s=doc.styles.add_style(name,1)
    s.font.name='Times New Roman';s.font.size=Pt(13);s.font.color.rgb=RGBColor.from_string(BLACK)
    fonts=s.element.get_or_add_rPr().rFonts
    for k in list(fonts.attrib):
        if k.endswith('Theme'):del fonts.attrib[k]
    for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Times New Roman')
    for border in s.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.3
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.widow_control=True
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for level,size in [(1,16),(2,14),(3,13)]:
    s=doc.styles[f'Heading {level}'];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(12 if level>1 else 0)
    s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if level==1:s.paragraph_format.page_break_before=True
doc.styles['Caption'].font.size=Pt(12)
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
settings=doc.settings.element
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
doc.core_properties.author='Kiều Trường Giang'
doc.core_properties.title='Tiểu luận 03 AI ML CNN RNN Knowledge Graph và RAG'
doc.core_properties.subject='Phát triển các hệ thống thông minh'
chapter=0;fig_count=Counter();table_count=Counter();eq_count=Counter();provenance=[];headings=[]


def clean(t):
    t=re.sub(r'\[\d+(?:\s*,\s*\d+)*\]','',t)
    t=re.sub(r'\s+([,.;:])',r'\1',t)
    t=re.sub(r'\s+',' ',t).strip()
    t=unicodedata.normalize('NFKC',t.replace('10⁻⁵','0,00001')).replace('\u00ad','')
    if '•' in t and not t.startswith('•'):
        t='• '+re.sub(r'\s*•\s*',' ',t).strip()
    t=re.sub(r'(?<=[A-Za-z])-\s+(?=[a-z])','-',t)
    return t.replace(':.','.')


def run(p,text,color=BLACK,bold=False):
    r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(13)
    r.font.color.rgb=RGBColor.from_string(color);r.bold=bold
    return r


def para(text,color=BLACK,source=None,bold=False,original=None):
    text=clean(text)
    if not text:return None
    p=doc.add_paragraph()
    if original is not None and color==RED:
        old=re.findall(r'\S+\s*',clean(original));new=re.findall(r'\S+\s*',text)
        matcher=SequenceMatcher(a=old,b=new,autojunk=False)
        if matcher.ratio()>.45:
            for tag,a,b,c,d in matcher.get_opcodes():
                if c!=d:run(p,''.join(new[c:d]),BLACK if tag=='equal' else RED,bold)
        else:run(p,text,color,bold)
    else:run(p,text,color,bold)
    if text.startswith(('• ','→ ')):
        p.paragraph_format.left_indent=Cm(.4);p.paragraph_format.first_line_indent=Cm(-.3)
    provenance.append({'kind':'paragraph','chapter':chapter,'color':color,'source':source,'text':text})
    return p


def heading(text,level=2,color=BLACK):
    p=doc.add_paragraph(clean(text),f'Heading {level}')
    for r in p.runs:r.font.color.rgb=RGBColor.from_string(color)
    headings.append((level,clean(text)))
    provenance.append({'kind':'heading','chapter':chapter,'color':color,'text':clean(text)})
    return p


def start_chapter(n,title):
    global chapter
    chapter=n;heading(f'CHƯƠNG {n}. {title}',1)


def link_line(label,url,color=RED):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after=Pt(6);run(p,label+': ',color)
    relation=doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),relation)
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    for tag,value in [('color',color),('sz','24')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);pr.append(el)
    rr.append(pr);text=OxmlElement('w:t');text.text=url;rr.append(text);link.append(rr);p._p.append(link)
    provenance.append({'kind':'link','chapter':chapter,'color':color,'label':label,'url':url})
    return p


def table(headers,rows,caption,color=BLACK,widths=None):
    table_count[chapter]+=1
    para(f'Bảng {chapter}.{table_count[chapter]}. {caption}',color).paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[16/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for rowvals in [headers]+list(rows):
        cells=t.rows[0].cells if rowvals is headers else t.add_row().cells
        for i,(c,value) in enumerate(zip(cells,rowvals)):
            c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            p=c.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(3)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 or len(str(value))>25 else WD_ALIGN_PARAGRAPH.CENTER
            rr=run(p,clean(str(value)),color,bold=rowvals is headers);rr.font.size=Pt(11.5)
            if rowvals is headers:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');tcpr.append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    provenance.append({'kind':'table','chapter':chapter,'color':color,'caption':caption,'rows':len(rows)})
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def picture(path,caption,color=BLACK,max_width=15.5,max_height=9.2,source=None):
    path=ROOT/path if not Path(path).is_absolute() else Path(path)
    if not path.exists():raise FileNotFoundError(path)
    im=Image.open(path);w,h=im.size;width=min(max_width,max_height*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    fig_count[chapter]+=1
    c=doc.add_paragraph(style='Caption');run(c,f'Hình {chapter}.{fig_count[chapter]}. {caption}',color)
    c.paragraph_format.keep_with_next=False
    provenance.append({'kind':'figure','chapter':chapter,'color':color,'source':source,'path':str(path.relative_to(ROOT)),'caption':caption})


def equation(latex,color=BLACK):
    eq_count[chapter]+=1
    data=subprocess.run(['pandoc','-f','markdown','-t','docx'],input=('$$'+latex+'$$').encode(),stdout=subprocess.PIPE,stderr=subprocess.PIPE,check=True).stdout
    xml=etree.fromstring(ZipFile(BytesIO(data)).read('word/document.xml'))
    mathnode=xml.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
    if mathnode is None:raise ValueError(latex)
    mathnode=deepcopy(mathnode)
    for mr in mathnode.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}r'):
        for old in list(mr.findall(qn('w:rPr'))):mr.remove(old)
        rp=OxmlElement('w:rPr');cl=OxmlElement('w:color');cl.set(qn('w:val'),color);rp.append(cl);mr.insert(0,rp)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=False
    p._p.append(mathnode);run(p,f'    ({chapter}.{eq_count[chapter]})',color)
    provenance.append({'kind':'equation','chapter':chapter,'color':color,'latex':latex})


# Trang bìa và mục lục là thay đổi trình bày, giữ màu đen.
for text,size,bold in [('HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG',14,True),('KHOA CÔNG NGHỆ THÔNG TIN 1',14,True)]:
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=run(p,text,BLACK,bold);r.font.size=Pt(size)
doc.add_paragraph().paragraph_format.space_after=Pt(40)
p=doc.add_paragraph('TIỂU LUẬN 03','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size=Pt(25);p.runs[0].bold=True
for text in ['TRÍ TUỆ NHÂN TẠO VÀ HỌC MÁY','CNN, RNN, KNOWLEDGE GRAPH VÀ RAG']:
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=run(p,text,BLACK,True);r.font.size=Pt(17)
doc.add_paragraph().paragraph_format.space_after=Pt(30)
for text in ['Học phần: Phát triển các hệ thống thông minh','Giảng viên hướng dẫn: Thầy Trần Đình Quế','Sinh viên thực hiện: Kiều Trường Giang','Mã sinh viên: B23DCCN254']:
    p=para(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10)
# Không tự suy ra tên lớp: cả hai bìa PDF không ghi trường này.
doc.add_paragraph().paragraph_format.space_after=Pt(45)
p=para('Hà Nội - 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break();p=doc.add_paragraph('MỤC LỤC','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-3" \\h \\z \\u '
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
for el in [begin,instruction,sep,end]:p.add_run()._r.append(el)


EQMAP={
 'tl01_p24_img23.png':r'w_{t+1}=w_t-\eta\nabla\mathcal{L}(w_t)',
 'tl01_p24_img24.png':r'\hat y=w^T x+b=\sum_{j=1}^{d}w_jx_j+b',
 'tl01_p24_img25.png':r'\mathcal{L}=\frac{1}{N}\sum_{i=1}^{N}(\hat y_i-y_i)^2',
 'tl01_p24_img26.png':r'\frac{\partial\mathcal{L}}{\partial w}=\frac{2}{N}X^T(\hat y-y)',
 'tl01_p24_img27.png':r'w\leftarrow w-\eta\frac{\partial\mathcal{L}}{\partial w}',
 'tl01_p25_img25.png':r'\sigma(z)=\frac{1}{1+e^{-z}}',
 'tl01_p25_img26.png':r'\mathcal{L}_{BCE}=-\frac{1}{N}\sum_{i=1}^{N}[y_i\log\hat y_i+(1-y_i)\log(1-\hat y_i)]',
 'tl01_p25_img27.png':r'p_k=\frac{e^{z_k}}{\sum_{j=1}^{K}e^{z_j}}',
 'tl01_p25_img28.png':r'\mathcal{L}_{CE}=-\frac{1}{N}\sum_{i=1}^{N}\sum_{k=1}^{K}y_{ik}\log p_{ik}',
 'tl01_p27_img32.png':r'E[(y-\hat f(x))^2]=\operatorname{Bias}[\hat f(x)]^2+\operatorname{Var}[\hat f(x)]+\sigma^2',
 'tl01_p28_img25.png':r'z^{(l)}=W^{(l)}a^{(l-1)}+b^{(l)},\quad a^{(l)}=\operatorname{ReLU}(z^{(l)})',
 'tl01_p28_img26.png':r'\frac{\partial\mathcal{L}}{\partial W^{(l)}}=\delta^{(l)}(a^{(l-1)})^T',
 'tl01_p28_img27.png':r'\delta^{(l)}=((W^{(l+1)})^T\delta^{(l+1)})\odot\operatorname{ReLU}^{\prime}(z^{(l)})',
 'tl01_p37_img24.png':r'Z(i,j)=\sum_{u=0}^{k_H-1}\sum_{v=0}^{k_W-1}X(i+u,j+v)K(u,v)+b',
 'tl01_p37_img25.png':r'H_{\mathrm{out}}=\left\lfloor\frac{H-k_H+2P}{S}\right\rfloor+1',
 'tl01_p37_img26.png':r'P_{\max}(i,j)=\max_{(u,v)\in R_{ij}}Z(u,v)',
 'tl01_p45_img24.png':r'\begin{gathered}h_t=\tanh(W_xx_t+W_hh_{t-1}+b)\\\hat y_t=W_yh_t+b_y\end{gathered}',
 'tl01_p45_img25.png':r'\frac{\partial\mathcal L}{\partial\theta}=\sum_{t=1}^{T}\sum_{k=1}^{t}\frac{\partial\mathcal L_t}{\partial h_t}\frac{\partial h_t}{\partial h_k}\left(\frac{\partial h_k}{\partial\theta}\right)_{\mathrm{direct}}',
 'tl01_p46_img20.png':r'f_t=\sigma(W_f[h_{t-1},x_t]+b_f)',
 'tl01_p46_img21.png':r'i_t=\sigma(W_i[h_{t-1},x_t]+b_i),\quad\tilde c_t=\tanh(W_c[h_{t-1},x_t]+b_c)',
 'tl01_p46_img22.png':r'c_t=f_t\odot c_{t-1}+i_t\odot\tilde c_t',
 'tl01_p46_img23.png':r'o_t=\sigma(W_o[h_{t-1},x_t]+b_o),\quad h_t=o_t\odot\tanh(c_t)',
 'tl01_p46_img24.png':r'\begin{gathered}z_t=\sigma(W_z[h_{t-1},x_t]+b_z)\\r_t=\sigma(W_r[h_{t-1},x_t]+b_r)\\\tilde h_t=\tanh(W[r_t\odot h_{t-1},x_t]+b)\\h_t=(1-z_t)\odot h_{t-1}+z_t\odot\tilde h_t\end{gathered}',
}


def theory_page(pg,lower=0,upper=1000):
    """Ghép lại các dòng trong cùng đoạn, giữ văn bản nguồn và công thức."""
    page=S1[pg-1];events=[]
    for b in page['blocks']:
        if b['bbox'][1]<lower or b['bbox'][1]>=upper:continue
        if b['type']=='image':events.append((b['bbox'][1],'image',b));continue
        for l in b['lines']:
            if l['bbox'][1]<lower or l['bbox'][1]>=upper:continue
            # Các bảng gốc được dựng lại dưới dạng Word editable ở bên dưới.
            if any(t['bbox'][1]-2<=l['bbox'][1]<=t['bbox'][3]+2 for t in page['tables'] if pg!=28):continue
            events.append((l['bbox'][1],'line',l))
    events.sort(key=lambda e:e[0]);pending=[];previous=None
    def flush():
        nonlocal pending
        if not pending:return
        text=clean(' '.join(pending));pending=[]
        if not text:return
        source=f'TL01 trang {pg}'
        if text.startswith('2. Thu nhỏ trọng số về gần 0'):return
        # Chuẩn hóa số mục, không phải cải tiến kiến thức.
        if re.match(r'^\d+\.\d+(?:\.\d+)?\.\s',text):
            if text=='1.7. Tiểu kết chương 1':text='1.8. Tiểu kết chương 1'
            level=3 if re.match(r'^\d+\.\d+\.\d+\.',text) else 2
            heading(text,level);return
        color=BLACK;original=text
        neutral={
          'Lịch sử phát triển gần 8 thập kỷ':'Lịch sử phát triển AI bao gồm các giai đoạn kỳ vọng cao và những đợt suy giảm đầu tư. Nhiều thay đổi về phương pháp xuất phát từ giới hạn kỹ thuật của giai đoạn trước, cùng sự xuất hiện của dữ liệu và phần cứng mới.',
          'Lịch sử phát triển của Trí tuệ Nhân tạo là một hành trình':'Lịch sử AI thể hiện sự chuyển đổi giữa nhiều hướng tiếp cận: từ các nền tảng lý thuyết của Alan Turing và Perceptron, qua những giai đoạn suy giảm đầu tư, đến học máy thống kê và học sâu. Các bước tiến gắn với thay đổi về dữ liệu, năng lực tính toán và phương pháp học.',
          'Những chu kỳ thăng trầm lịch sử đã để lại':'Không có một mô hình đơn lẻ phù hợp với mọi bài toán. Việc lựa chọn cần dựa trên cấu trúc dữ liệu, mục tiêu đánh giá và khả năng kiểm tra kết quả, từ các kỹ thuật học máy cho bảng thuộc tính đến kiến trúc chuyên biệt cho ảnh và chuỗi.',
          'Đó chính là tiền đề dẫn dắt chúng ta':'Chương tiếp theo cụ thể hóa các nguyên tắc trên bằng hai bài toán dữ liệu bảng, trong đó phép chia dữ liệu và lựa chọn thước đo được xác định trước khi so sánh mô hình.'}
        for prefix,replacement in neutral.items():
            if text.startswith(prefix):text=replacement;color=RED
        replacements={
          'cực kỳ':'rất',
          'Cực kỳ':'Rất',
          'Cân bằng hoàn hảo':'Cân bằng',
          'tiêu chuẩn vàng':'cách làm phổ biến',
          'Triệt tiêu phương sai':'Giảm phương sai',
          'triệt tiêu phương sai':'giảm phương sai',
          'giảm sốc':'giảm',
          'tạo cơn địa chấn toàn cầu':'thu hút sự chú ý',
          'ngọn lửa cách mạng học sâu':'sự phát triển của học sâu',
          'hạ gục kỳ thủ':'thắng kỳ thủ',
          'thống trị tuyệt đối của Deep Learning trước các phương pháp kỹ nghệ đặc trưng thủ công':'ưu thế của mạng sâu trên benchmark này so với các phương pháp thiết kế đặc trưng thủ công',
          'bước tiến vượt bậc':'bước phát triển',
          'khả năng thích ứng phi thường':'khả năng thích ứng',
          'π(a\\s)':'π(a|s)',
          'LSTM), giải quyết triệt để vấn đề biến mất gradient':'LSTM), giảm khó khăn học phụ thuộc dài hạn do biến mất gradient',
          'chính thức đặt tên cho thuật ngữ "Deep Learning"':'góp phần khôi phục quan tâm đến huấn luyện mạng sâu',
          'Mỗi bước tiến công nghệ đều bắt nguồn':'Nhiều bước tiến công nghệ bắt nguồn',
          'Hồi quy (Regression): Khi nhãn y ∈ {C1, C2, ..., Ck} là giá trị số liên tục':'Hồi quy (Regression): Khi nhãn y là giá trị số liên tục',
          '3. Translation invariance: Phát hiện đặc trưng bất kể vị trí':'3. Translation equivariance: phép tích chập dùng chung kernel khi dịch đầu vào; pooling tạo bất biến cục bộ ở mức độ nhất định.',
        }
        if text.startswith('• Hồi quy (Regression):'):
            text='• Hồi quy (Regression): Khi nhãn y thuộc tập số thực, đầu ra là giá trị liên tục (ví dụ: dự báo giá nhà, nhiệt độ, doanh thu).';color=RED
        for old,new in replacements.items():
            if old in text:text=text.replace(old,new);color=RED
        if '2006' in text and 'đặt tên' in text:text=text.replace('chính thức đặt tên cho thuật ngữ "Deep Learning"','góp phần khôi phục quan tâm đến huấn luyện mạng sâu');color=RED
        if 'giải quyết triệt để' in text:text=text.replace('giải quyết triệt để','giảm bớt');color=RED
        if 'Trong chương này, ta triển khai một kiến trúc' in text:text=text.replace('Trong chương này, ta triển khai','Có thể minh họa lý thuyết bằng');color=RED
        if 'Toàn bộ cơ chế' in text:text='Trong đó δ là sai số lan truyền ngược. Quy tắc chuỗi cho phép truyền gradient qua các lớp; phần thực nghiệm mới sử dụng các mô hình được mô tả riêng bên dưới.';color=RED
        if 'Với bài toán phân loại (Diabetes Risk)' in text:text=text.replace('(Diabetes Risk)','đa lớp');color=RED
        if text.startswith('Dense layer và CNN xử lý từng mẫu độc lập'):
            text='Dense layer không giữ trạng thái qua thời gian. CNN có thể xử lý tín hiệu 1D theo trường tiếp nhận hữu hạn; RNN bổ sung trạng thái truyền qua các bước để biểu diễn ngữ cảnh của chuỗi.';color=RED
        if text.startswith('• Wx, Wh, Wi:'):text='• Wₓ, Wₕ và Wᵧ là ma trận trọng số; các tham số được dùng chung tại mọi bước thời gian.';color=RED
        if text.startswith('Năm 1956, tại hội thảo lịch sử Dartmouth'):
            text='Hội thảo Dartmouth năm 1956 đánh dấu sự hình thành của AI như một lĩnh vực nghiên cứu. John McCarthy dùng tên gọi Artificial Intelligence cho hướng xây dựng các hệ thống máy tính thông minh.';color=RED
        if text.startswith('• 1936') and 'mọi quy trình tư duy' in text:
            text=text.replace('mọi quy trình tư duy có thể biểu diễn hữu hạn đều có thể được máy móc tính toán.','khái niệm tính được bằng thuật toán có thể được mô hình hóa chặt chẽ; điều này không chứng minh mọi hoạt động tư duy đều tính được.');color=RED
        if 'hàng nghìn megawatt điện' in text:text=text.replace('hàng nghìn megawatt điện','năng lượng đáng kể, phụ thuộc quy mô và phần cứng');color=RED
        if 'phân loại' in text and 'ngày nay' in text:pass
        # Mô tả hiện đại không liệt kê phiên bản sản phẩm như các mốc 2026 đã được xác minh.
        if text.startswith('• 2023–2026') or text.startswith('• 2023-2026'):
            text='• Từ 2023: Mô hình nền tảng đa phương thức kết hợp văn bản, hình ảnh và âm thanh. Các tác tử dùng mô hình ngôn ngữ để lựa chọn công cụ và thực hiện chuỗi thao tác, đồng thời cần kiểm soát quyền, đánh giá lỗi và giám sát con người.';color=RED
        if 'Standard Scaling' in text:
            text='• Chuẩn hóa thang đo: Đưa các thuộc tính về cùng một miền giá trị bằng Standard Scaling hoặc Min-Max Scaling nhằm giúp các thuật toán hạ gradient hội tụ nhanh và ổn định.'
            para(text,BLACK,source);equation(r'z=\frac{x-\mu}{\sigma},\qquad x^{\prime}=\frac{x-x_{\min}}{x_{\max}-x_{\min}}')
            para('x là giá trị gốc; μ, σ, x_min và x_max lần lượt là trung bình, độ lệch chuẩn và hai biên của đặc trưng. Các đại lượng này chỉ được ước lượng từ train, rồi giữ cố định khi biến đổi validation/test.',RED)
            return
        if text.startswith('• Dữ liệu huấn luyện') and 'nhãn mục tiêu' in text:
            para('• Dữ liệu huấn luyện bao gồm cặp mẫu đầu vào và nhãn mục tiêu tương ứng:',BLACK,source)
            equation(r'D=\{(x_i,y_i)\}_{i=1}^{N}')
            return
        if text.startswith('• Dữ liệu huấn luyện') and 'không có nhãn' in text:
            para('• Dữ liệu huấn luyện chỉ bao gồm các mẫu đầu vào mà không có nhãn đích:',BLACK,source)
            equation(r'D=\{x_i\}_{i=1}^{N}')
            return
        if text.startswith('• Mục tiêu: Tìm hàm ánh xạ'):
            para('• Mục tiêu: Tìm hàm ánh xạ từ không gian đầu vào sang không gian đầu ra sao cho dự báo xấp xỉ nhãn thực tế.',BLACK,source)
            equation(r'f:\mathcal{X}\rightarrow\mathcal{Y},\qquad \hat y=f(x)')
            return
        if text.startswith('• Điều chuẩn L2'):
            para('• Điều chuẩn L2 (Ridge Regression / Weight Decay): Phạt tổng bình phương trọng số. Thu nhỏ trọng số về gần 0 nhưng không triệt tiêu hoàn toàn.',BLACK,source)
            equation(r'\mathcal{L}_{reg}=\mathcal{L}+\lambda\sum_i w_i^2')
            return
        if text.startswith('• Điều chuẩn L1'):
            para('• Điều chuẩn L1 (Lasso Regression): Phạt tổng giá trị tuyệt đối của trọng số. Có khả năng triệt tiêu trọng số về chính xác bằng 0, đóng vai trò chọn lọc đặc trưng tự động.',BLACK,source)
            equation(r'\mathcal{L}_{reg}=\mathcal{L}+\lambda\sum_i|w_i|')
            return
        if 'Dành cho bài toán phân loại' in text:pass
        if 'Tiêu chuẩn phân chia tối ưu' in text:
            para('• Tiêu chuẩn phân chia tối ưu: sử dụng độ lợi thông tin (Information Gain) dựa trên Entropy hoặc chỉ số Gini Impurity.',RED,source,original=original)
            equation(r'H(S)=-\sum_{k=1}^{K}p_k\log_2p_k,\qquad G(S)=1-\sum_{k=1}^{K}p_k^2',RED)
            para('pₖ là tỷ lệ mẫu thuộc lớp k ở nút đang xét. Độ lợi thông tin là độ giảm entropy có trọng số sau khi chia. Một phép chia có độ lợi cao tạo các nút con đồng nhất hơn về nhãn.',RED)
            return
        if text.startswith('Bảng so sánh tổng hợp các thuật toán'):
            t=table(['Thuật toán','Bản chất','Độ phức tạp','Giải thích','Ưu điểm','Hạn chế'],[
              ['Logistic Regression','Tuyến tính + Sigmoid','Dự đoán O(d)','Hệ số trọng số','Đơn giản, nhanh; xuất xác suất','Biên phân cách tuyến tính'],
              ['Decision Tree','Cắt lát không gian','Phụ thuộc độ sâu','Cây luật NẾU–THÌ','Trực quan; ít cần scaling','Dễ quá khớp, nhạy nhiễu'],
              ['Random Forest','Bagging nhiều cây','Phụ thuộc số cây, N, d và độ sâu','Feature importance','Bền bỉ, giảm phương sai','Mô hình lớn hơn, dự báo chậm hơn'],
              ['Gradient Boosting','Boosting sửa sai số','Phụ thuộc số vòng và độ sâu','Trung bình','Mạnh với dữ liệu bảng','Nhạy siêu tham số, tốn thời gian'],
              ['SVM','Cực đại hóa margin','Phụ thuộc kernel và số mẫu','Thấp với kernel','Mạnh với ít mẫu, nhiều chiều','Chậm trên tập lớn'],
              ['KNN','Khoảng cách mẫu','Brute-force O(Nd) mỗi truy vấn','Láng giềng','Trực quan, không giả định phân phối','Tốn bộ nhớ, dự báo chậm']], 'So sánh các thuật toán học máy kinh điển',widths=[2.6,2.3,2.4,3.1,2.8,2.8])
            for row,col in [(2,4),(3,2),(4,2),(5,2),(6,2)]:
                for rr in t.cell(row,col).paragraphs[0].runs:rr.font.color.rgb=RGBColor.from_string(RED)
            return
        para(text,color,source,original=original if color==RED else None)
    for y,typ,b in events:
        if typ=='image':
            flush();filename=Path(b['path']).name
            if filename in EQMAP:
                equation(EQMAP[filename],RED if filename=='tl01_p45_img25.png' else BLACK)
                if filename=='tl01_p45_img25.png':para('θ là một tham số dùng chung của mạng; đạo hàm cuối chỉ lấy tác động trực tiếp tại bước k. Tổng theo k cộng các đường mà tham số ảnh hưởng tới loss ở bước t. Đạo hàm giữa hai trạng thái bao gồm toàn bộ các bước trung gian, với đạo hàm đồng nhất khi k=t.',RED)
            elif pg in [12,17]:picture(b['path'],'Các giai đoạn và quan hệ giữa các kỹ thuật AI',source=f'TL01 trang {pg}')
            else:
                # Các phương trình CNN/RNN gốc giữ nguyên hình nếu chưa ánh xạ OMML.
                eq_count[chapter]+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                im=Image.open(ROOT/b['path']);w,h=im.size
                p.add_run().add_picture(str(ROOT/b['path']),width=Cm(min(13.,(b['bbox'][2]-b['bbox'][0])/72*2.54)))
                run(p,f'    ({chapter}.{eq_count[chapter]})')
            previous=None;continue
        t=b['text'];ishead=bool(re.match(r'^\d+\.\d+(?:\.\d+)?\.\s',t))
        if t.startswith('CHƯƠNG'):continue
        if previous is not None:
            gap=b['bbox'][1]-previous['bbox'][3]
            if gap>8 or ishead or t.startswith(('• ','1. ','2. ','3. ','4. ','Giai đoạn')):flush()
        if ishead:flush()
        pending.append(t);previous=b
    flush()


start_chapter(1,'LỊCH SỬ VÀ PHÁT TRIỂN AI')
add_intro(1,globals())
theory_page(11,100,405)
table(['Chiều kích','Tiếp cận dựa trên con người','Tiếp cận dựa trên tính hợp lý'],[
 ['Tư duy','Tư duy như con người (Thinking Humanly). Khoa học nhận thức; mô phỏng quá trình nơ-ron não bộ xử lý thông tin.','Tư duy duy lý (Thinking Rationally). Dựa trên các quy luật logic hình thức (Boole, Frege); rút ra kết luận đúng đắn từ các tiền đề.'],
 ['Hành động','Hành động như con người (Acting Humanly). Tiêu biểu là phép thử Turing (1950); hành vi máy khiến con người tin rằng nó là người.','Hành động duy lý (Acting Rationally). Tác tử hợp lý quan sát môi trường và hành động tối đa hóa kỳ vọng đạt mục tiêu.']], 'Bốn hướng tiếp cận trí tuệ nhân tạo',widths=[2.5,6.75,6.75])
theory_page(11,lower=736)
for pg in range(12,18):theory_page(pg,upper=493 if pg==17 else 1000)
table(['Trụ cột','Bản chất kỹ thuật','Cột mốc minh họa'],[
 ['Dữ liệu số khổng lồ','Mạng nơ-ron sâu chứa hàng triệu đến hàng tỷ tham số, do đó đòi hỏi tập dữ liệu khổng lồ để tránh overfitting và đạt khả năng tổng quát hóa cao.','Dự án ImageNet gán nhãn hơn 14 triệu hình ảnh, cung cấp dữ liệu tiêu chuẩn cho mô hình học sâu.'],
 ['Tính toán GPU','Mạng nơ-ron là chuỗi phép nhân ma trận song song; kiến trúc nhiều nhân của GPU phù hợp hơn CPU truyền thống.','CUDA xuất hiện năm 2007; AlexNet năm 2012 huấn luyện trên hai GPU GTX 580.'],
 ['Thuật toán và kiến trúc','ReLU, Dropout, BatchNorm, Adam và kết nối tắt hỗ trợ tối ưu mạng sâu và hạn chế quá khớp.','Dropout (2012), BatchNorm (2015), Adam (2014), ResNet (2015).']], 'Ba trụ cột phát triển của học sâu',widths=[3,6.5,6.5])
theory_page(18,lower=330)
theory_page(19);theory_page(20,upper=200)
heading('1.7. Mối liên hệ giữa các giai đoạn phát triển',2,RED)
para('AI biểu tượng cung cấp cách biểu diễn tri thức và suy luận có thể kiểm tra. Học máy thống kê chuyển trọng tâm sang việc học quy luật từ dữ liệu; học sâu tiếp tục học biểu diễn thay cho nhiều bước thiết kế đặc trưng thủ công. Transformer cho phép mô hình hóa quan hệ giữa các vị trí qua attention và huấn luyện song song hơn RNN. KG và RAG kết nối mô hình dữ liệu với tri thức tường minh; chúng kế thừa các ý tưởng cũ thay vì thay thế toàn bộ AI biểu tượng.',RED)
para('Trong báo cáo này, mối liên hệ đó được thể hiện bằng ba loại đầu vào: bảng thuộc tính, ảnh và chuỗi thời gian. Các mô hình tạo ra dự đoán định lượng; Chương 5 bổ sung bằng chứng và ngữ cảnh để người dùng hiểu phạm vi của các kết quả.',RED)
add_details('ai_bridge',globals())
theory_page(20,lower=200)

start_chapter(2,'CÁC KỸ THUẬT MACHINE LEARNING CƠ BẢN')
add_intro(2,globals())
for pg in range(21,29):theory_page(pg,lower=610 if pg==21 else 0,upper=609 if pg==28 else 1000)
heading('2.1.9. Kiểm định chéo và lựa chọn siêu tham số',3,RED)
para('Train dùng để ước lượng tham số; validation dùng để kiểm tra các lựa chọn phát triển và checkpoint; test chỉ dùng cho đánh giá cuối. Trong phân loại, StratifiedKFold giữ tỷ lệ lớp trong từng fold. Với giá nhà, GroupKFold giữ các bản ghi cùng địa chỉ trong một fold. Toàn bộ imputation, scaling và one-hot encoding được fit bên trong từng fold, không fit trên toàn bộ dữ liệu trước khi chia.',RED)
para('Grid search so sánh tập siêu tham số hữu hạn đã định trước. Dùng macro-F1 khi các lớp mất cân bằng và MAE theo tỷ VND khi dự báo giá nhà. Không lựa chọn mô hình bằng chỉ số test. Khoảng cách train–validation–test giúp phát hiện quá khớp; một lần chia và một random seed chưa mô tả hết độ bất định.',RED)
heading('2.1.10. Các thước đo đánh giá',3,RED)
equation(r'\operatorname{MAE}=\frac{1}{n}\sum_{i=1}^{n}|y_i-\hat y_i|,\quad \operatorname{MSE}=\frac{1}{n}\sum_{i=1}^{n}(y_i-\hat y_i)^2',RED)
equation(r'\operatorname{RMSE}=\sqrt{\operatorname{MSE}},\quad R^2=1-\frac{\sum_i(y_i-\hat y_i)^2}{\sum_i(y_i-\bar y)^2}',RED)
para('Trong (2.%d)–(2.%d), n là số mẫu, yᵢ là giá trị thực, ŷᵢ là dự đoán và ȳ là trung bình của nhãn trong tập đánh giá. MAE và RMSE có cùng đơn vị với nhãn, MSE có đơn vị bình phương. R² so với dự đoán hằng bằng trung bình, có thể âm và không phải phần trăm chính xác.'%(eq_count[2]-1,eq_count[2]),RED)
equation(r'\operatorname{Accuracy}=\frac{TP+TN}{TP+TN+FP+FN},\quad P=\frac{TP}{TP+FP},\quad R=\frac{TP}{TP+FN}',RED)
equation(r'F_1=\frac{2PR}{P+R},\qquad F_{1,\mathrm{macro}}=\frac{1}{K}\sum_{k=1}^{K}F_{1,k}',RED)
para('TP, TN, FP và FN lần lượt là số dự đoán dương đúng, âm đúng, dương sai và âm sai trong cách xét một lớp so với các lớp còn lại. Với K lớp, macro-F1 cho mỗi lớp trọng số bằng nhau, vì vậy thể hiện rõ hơn hiệu năng ở lớp ít mẫu.',RED)

d=TAB['diabetes_classifier'];h=TAB['housing']
heading('2.2. Phân tích và chuẩn bị dữ liệu thực nghiệm',2,RED)
table(['Dataset','Quy mô','Mục tiêu','Train / Val / Test'],[
 ['Diabetes Risk Prediction','50.000 dòng; 41 cột gồm cả nhãn','Ba nhãn High, Low, Moderate',f"{d['split']['train']} / {d['split']['validation']} / {d['split']['test']}"],
 ['Vietnam Housing 2024','30.229 dòng; 12 cột gốc','Giá nhà theo tỷ VND',f"{h['split']['train']} / {h['split']['validation']} / {h['split']['test']}"]], 'Dữ liệu của các thực nghiệm tự xây dựng lại',RED,[4.5,3.6,4.1,3.8])
para('Diabetes Risk Prediction là dữ liệu tổng hợp phục vụ học máy, không được coi là 50.000 hồ sơ bệnh nhân đã xác minh lâm sàng. Phân phối thực tế gồm High 36.593 mẫu (73,2%), Moderate 12.937 mẫu (25,9%) và Low 470 mẫu (0,9%). Mất cân bằng giữa ba lớp cần được tính đến khi chọn mô hình và đọc chỉ số. Tám đầu vào gồm Age, BMI, HbA1c, Fasting_Blood_Sugar, Blood_Glucose, Blood_Pressure_Systolic, HDL và LDL. Patient_ID, Diabetes_Risk_Score, AI_Health_Recommendation và Doctor_Consultation_Needed không được sử dụng vì là định danh hoặc có thể tiết lộ nhãn.',RED)
para('Dữ liệu giá nhà có các trường Area, Frontage, Access Road, Floors, Bedrooms, Bathrooms cùng hướng nhà, hướng ban công, pháp lý và nội thất. Province được trích từ phần cuối địa chỉ. Pipeline điền thiếu số bằng trung vị train, mã hóa one-hot với nhóm chưa gặp được bỏ qua, và học log(1 + Price) rồi biến đổi ngược trước khi tính sai số. Các bản ghi cùng địa chỉ được giữ trong cùng một tập. Đây là kiểm tra tổng quát hóa sang địa chỉ khác, chưa phải backtest giá nhà theo thời gian.',RED)
picture(FIG/'housing_eda.png','Phân phối giá nhà trước và sau biến đổi logarit',RED)
add_details('ml_data',globals())
heading('2.3. So sánh và lựa chọn mô hình',2,RED)
rows=list(csv.DictReader((DATA/'diabetes_cv.csv').open()))
table(['Mô hình','CV macro-F1'],[[r['model'],f"{float(r['CV_F1_macro']):.4f}"] for r in rows], 'So sánh phân loại với StratifiedKFold 3 phần',RED,[10,6])
para(f"Random Forest được chọn bằng CV macro-F1 {d['CV_F1_macro']:.4f}, gồm 80 cây, max_depth=10, class_weight=balanced, random_state=42. Macro-F1 train {d['train_metrics']['F1_macro']:.4f}, validation {d['validation_metrics']['F1_macro']:.4f} và test {d['metrics']['F1_macro']:.4f} cho thấy vẫn có khoảng cách khái quát hóa. Không dùng test để tăng độ sâu cây.",RED)
rows=list(csv.DictReader((DATA/'housing_cv.csv').open()))
table(['Mô hình','CV MAE (tỷ VND)'],[[r['model'],f"{float(r['CV_MAE']):.4f}"] for r in rows], 'So sánh hồi quy với GroupKFold 3 phần',RED,[10,6])
para('Hồi quy chọn Log Random Forest với 80 cây, max_depth=14, min_samples_leaf=3. Ridge thử alpha 1, 10, 100; Gradient Boosting thử learning_rate 0,04 và 0,08. Tất cả cấu hình cùng một nhóm train và các fold theo địa chỉ. Mô hình trung vị cung cấp mốc đánh giá độ hữu ích của các mô hình học được.',RED)
heading('2.4. Kết quả test độc lập và phân tích sai số',2,RED)
table(['Bài toán','Chỉ số','Kết quả'],[
 ['Phân loại nguy cơ','Accuracy / Precision macro',f"{d['metrics']['Accuracy']:.4f} / {d['metrics']['Precision_macro']:.4f}"],
 ['Phân loại nguy cơ','Recall macro / F1 macro',f"{d['metrics']['Recall_macro']:.4f} / {d['metrics']['F1_macro']:.4f}"],
 ['Giá nhà','MAE / RMSE (tỷ VND)',f"{h['metrics']['MAE']:.4f} / {h['metrics']['RMSE']:.4f}"],
 ['Giá nhà','MSE / R²',f"{h['metrics']['MSE']:.4f} / {h['metrics']['R2']:.4f}"]], 'Kết quả được tính từ predictions của main.ipynb',RED,[4,6,6])
picture(FIG/'diabetes_confusion.png','Ma trận nhầm lẫn của phân loại nguy cơ trên 10.000 mẫu test',RED,max_height=8)
para('Ma trận nhầm lẫn cho thấy Accuracy chịu ảnh hưởng lớn của lớp High. Macro-precision thấp hơn macro-recall phản ánh chi phí dự đoán nhầm khi dùng class_weight=balanced. Nhãn dự đoán chỉ mô tả quy tắc của dataset, không phải chẩn đoán hoặc nguy cơ bệnh ngoài đời.',RED)
picture(FIG/'housing_predictions.png','Giá nhà dự đoán và thực tế tại các địa chỉ giữ riêng',RED)
para(f"MAE train của giá nhà là {h['train_metrics']['MAE']:.4f} tỷ VND, validation {h['validation_metrics']['MAE']:.4f}, test {h['metrics']['MAE']:.4f}; R² test {h['metrics']['R2']:.4f}. Sai số tuyệt đối còn lớn đối với định giá cá nhân. Các thông tin vị trí trong địa chỉ, chất lượng tin đăng và thời điểm giao dịch chưa được mô hình hóa đầy đủ. Mức lệch giữa train và test cần được theo dõi trước khi tăng độ phức tạp.",RED)
add_details('ml_errors',globals())
heading('2.5. Kết luận chương',2)
para('Với dữ liệu dạng bảng, Random Forest và Gradient Boosting thường cạnh tranh tốt với Deep Learning. Không có một mô hình đơn lẻ phù hợp với mọi tập dữ liệu.',BLACK,'TL01 trang 36; diễn giải kết luận được giữ')
para('Hai bài toán cho thấy điểm số phụ thuộc cách xác định nhãn, phép chia dữ liệu và thước đo. Macro-F1 giúp nhận rõ hạn chế ở lớp hiếm; chia theo địa chỉ giúp đánh giá giá nhà ngoài các nhóm đã gặp. Dù Random Forest được chọn trong cả hai bài toán, chất lượng đầu vào và miền sử dụng vẫn là giới hạn chính. Các tệp checkpoint, dự đoán và bảng CV được giữ để có thể kiểm tra lại từng kết luận.',RED)

start_chapter(3,'MẠNG NƠ-RON TÍCH CHẬP CNN')
add_intro(3,globals())
theory_page(37,lower=220)
theory_page(38,upper=110)
table(['Tầng','Vai trò'],[['Conv2D','Trích xuất đặc trưng cục bộ: cạnh, góc, texture'],['BatchNorm','Chuẩn hóa giúp huấn luyện ổn định hơn'],['ReLU','Tạo tính phi tuyến'],['MaxPool','Giảm kích thước, tăng receptive field'],['Dropout','Điều chuẩn, hạn chế quá khớp'],['Dense và Softmax','Phân loại đa lớp']], 'Vai trò của các tầng CNN',widths=[4,12])
heading('3.1.5. Các kiến trúc CNN tiêu biểu',3)
table(['Kiến trúc','Năm','Đặc điểm'],[['LeNet-5','1998','Nhận dạng chữ số, convolution và subsampling'],['AlexNet','2012','ReLU, Dropout và huấn luyện GPU'],['VGG','2014','Các kernel 3 × 3 xếp chồng'],['GoogLeNet','2014','Các nhánh Inception'],['ResNet','2015','Kết nối residual'],['EfficientNet','2019','Compound scaling độ sâu, độ rộng và độ phân giải']], 'Các kiến trúc CNN và đặc trưng thiết kế đã hiệu chỉnh',RED,[3.5,2,10.5])
para('LeNet-5 (1998) hướng tới nhận dạng ký tự và chữ số; các kiến trúc về sau mở rộng khả năng học biểu diễn trên ảnh phức tạp hơn. Chỉ số Top-5 của các biến thể, dataset hoặc chế độ ensemble khác nhau không tạo thành một phép so sánh thống nhất. Bảng 3.2 đối chiếu đặc điểm kiến trúc; các kết quả định lượng của chương được tính riêng trên MNIST.',RED)
heading('3.1.6. Backpropagation, loss và optimizer',3,RED)
para('Lan truyền xuôi tạo logits của từng lớp; CrossEntropyLoss kết hợp log-softmax và negative log likelihood, vì vậy không áp softmax trước loss trong PyTorch. Backpropagation dùng quy tắc chuỗi để tính gradient của kernel và dense layer. Adam cập nhật trọng số dựa trên các ước lượng moment của gradient; learning rate, weight decay và checkpoint được lựa chọn mà không nhìn test.',RED)
equation(r'\frac{\partial\mathcal{L}}{\partial K_{uv}}=\sum_{i,j}\frac{\partial\mathcal{L}}{\partial Y_{ij}}X_{i+u,j+v}',RED)
para('Trong (3.%d), K là kernel, X là đầu vào và Y là đầu ra trước kích hoạt; chỉ số u,v xác định phần tử kernel và i,j chạy qua vị trí đầu ra. Cùng một trọng số được dùng tại nhiều vị trí nên gradient là tổng đóng góp ở các vị trí đó.'%eq_count[3],RED)
heading('3.1.7. Transfer learning và data augmentation',3,RED)
para('Transfer learning khởi tạo từ mạng đã học trên tập lớn, thay đầu phân loại rồi đóng băng hoặc fine-tune một số tầng. Nó hữu ích khi dữ liệu đích ít nhưng phải công bố phần pretrained và phần được học lại. Thực nghiệm MNIST dưới đây không dùng pretrained; augmentation chỉ dịch ảnh train ±1 pixel, giữ nguyên validation và test. Không dùng phép xoay mạnh làm thay đổi ý nghĩa giữa các chữ số.',RED)
heading('3.2. Dữ liệu ảnh và kiến trúc thực nghiệm',2,RED)
para('MNIST có 60.000 ảnh train gốc và 10.000 ảnh test độc lập, ảnh xám 28 × 28, nhãn 0–9. Tách 12.000 ảnh từ train làm validation theo stratification, còn 48.000 ảnh huấn luyện. Mỗi pixel được chia 255; không học thống kê từ test. Bộ tải đọc trực tiếp các tệp IDX nén gzip, kiểm tra magic number, số lượng ảnh và kích thước.',RED)
table(['Tầng','Kích thước đầu ra','Tham số'],[['Ảnh vào','1 × 28 × 28','Pixel / 255'],['Conv1 + ReLU + MaxPool','8 × 14 × 14','Kernel 3; padding 1; pool 2'],['Conv2 + ReLU + MaxPool','16 × 7 × 7','Kernel 3; padding 1; pool 2'],['Flatten + Dense + ReLU','64','784 → 64'],['Dense logits','10','64 → 10']], 'Kiến trúc CNN tự huấn luyện',RED,[5.5,4,6.5])
para(f"Mạng có {CNN['parameters']:,} tham số, train từ đầu {CNN['epochs']} epoch trên CPU, Adam learning_rate=0,001, batch_size=128, weight_decay=10⁻⁴ và gradient clipping norm=1. Checkpoint có validation loss nhỏ nhất nằm ở epoch {CNN['best_epoch']}. Thời gian thực nghiệm ghi trong metadata là {CNN['seconds']:.1f} giây trên môi trường hiện tại.",RED)
add_details('cnn_mechanism',globals())
heading('3.3. Kết quả và phân tích dự đoán',2,RED)
table(['Tập','Accuracy','Macro-F1'],[['Validation',f"{CNN['validation_metrics']['Accuracy']:.4f}",f"{CNN['validation_metrics']['F1_macro']:.4f}"],['Test 10.000 ảnh',f"{CNN['metrics']['Accuracy']:.4f}",f"{CNN['metrics']['F1_macro']:.4f}"]], 'Đánh giá CNN trên các tập độc lập',RED,[6,5,5])
picture(FIG/'cnn_learning.png','Learning curve của CNN qua 8 epoch',RED,max_height=7)
picture(FIG/'cnn_confusion.png','Ma trận nhầm lẫn CNN trên tập test MNIST',RED,max_height=10)
picture(FIG/'cnn_examples.png','Sáu dự đoán đầu và sáu lỗi đầu tiên theo thứ tự test',RED,max_height=7)
para(f"Accuracy test là {100*CNN['metrics']['Accuracy']:.2f}%, tương đương {int(round((1-CNN['metrics']['Accuracy'])*10000))} lỗi trên 10.000 ảnh. Các ảnh sai trong Hình 3.{fig_count[3]} được chọn bằng quy tắc cố định, không chọn riêng ví dụ đẹp. Learning curve cho thấy validation loss giảm trong cấu hình đã chạy; không kết luận mọi CNN hoặc mọi dữ liệu ảnh đều đạt mức này. Ảnh vẽ tay trên web có thể khác phân phối MNIST nên hiển thị confidence của mô hình thay cho cam kết đúng.",RED)
add_details('cnn_errors',globals())
heading('3.4. Suy luận, triển khai và phạm vi mở rộng',2,RED)
para('Checkpoint PyTorch được xuất thành trọng số JSON. API thực hiện convolution, ReLU, max-pooling, flatten và dense bằng đúng trọng số đã huấn luyện, dùng Python standard library để tránh đóng gói PyTorch nặng vào serverless. Bộ vector đối chiếu kiểm tra logits của API với PyTorch. Giao diện hỗ trợ ảnh test thật, canvas vẽ số, kiểm tra ảnh trống, trạng thái đang xử lý và thông báo lỗi.',RED)
para('ASL được giữ như ví dụ mở rộng của bài toán thị giác: các cử chỉ M/N hoặc U/V gần nhau, nền ảnh và góc chụp có thể làm giảm khả năng tổng quát hóa. Trong lần thực hiện này chưa huấn luyện lại ASL, vì vậy các accuracy và biểu đồ ASL cũ không được dùng như kết quả mới. Tương tự, không khẳng định Keras và PyTorch tương đương từ một phép thử chỉ chạy PyTorch.',RED)

start_chapter(4,'MẠNG NƠ-RON HỒI TIẾP RNN')
add_intro(4,globals())
theory_page(45,lower=297)
theory_page(46,upper=478)
table(['Mạng','Trạng thái','Số cổng','Đặc điểm'],[['RNN','Hidden state hₜ','Không có cổng','Ít tham số, khó học phụ thuộc xa'],['LSTM','hₜ và cell state cₜ','Input, forget, output','Luồng cộng qua cell state hỗ trợ truyền gradient'],['GRU','Hidden state hₜ','Reset và update','Ít tham số hơn LSTM cùng kích thước ẩn']], 'So sánh cấu trúc recurrent network',RED,[2,4,4,6])
heading('4.1.5. Cửa sổ trượt và các kiểu đầu vào đầu ra',3)
para('Cửa sổ trượt tạo một mẫu bằng L quan sát quá khứ và nhãn của bước kế tiếp. Many-to-one ánh xạ chuỗi thành một giá trị hoặc nhãn; one-to-many sinh chuỗi từ một ngữ cảnh đầu; many-to-many tạo một đầu ra ở mỗi bước hoặc một chuỗi đầu ra khác độ dài. Thực nghiệm dự báo giá trong chương sử dụng many-to-one.',RED)
equation(r'X_t=[r_{t-L+1},\ldots,r_t],\quad r_t=\log(P_t/P_{t-1}),\quad \hat P_{t+1}=P_t e^{\hat r_{t+1}}',RED)
para('Pₜ là giá đóng cửa, rₜ là log-return, L=30 là chiều dài cửa sổ. Dự báo tại t dùng dữ liệu không muộn hơn t; nhãn ở t+1. Mỗi cửa sổ khởi tạo trạng thái ẩn bằng 0 và không truyền hidden state từ test về train.',RED)
heading('4.1.6. BPTT và ổn định gradient',3,RED)
equation(r'\frac{\partial h_t}{\partial h_{t-k}}=\prod_{j=t-k+1}^{t}\frac{\partial h_j}{\partial h_{j-1}}',RED)
para('BPTT trải mạng theo thời gian và cộng gradient của cùng tham số qua các bước. Tích nhiều Jacobian có chuẩn nhỏ hơn 1 có thể làm gradient suy giảm, còn chuẩn lớn có thể gây bùng nổ. LSTM/GRU tạo đường truyền có cổng để giảm khó khăn này nhưng không bảo đảm giải quyết mọi phụ thuộc dài hạn. Gradient clipping giới hạn norm; chuẩn hóa chỉ trên train và chọn checkpoint bằng validation giúp ổn định thực nghiệm.',RED)
heading('4.2. Dữ liệu, phép chia và cấu hình',2,RED)
para('Dùng snapshot Yahoo Finance của Tiểu luận 02: FPT.VN có 1.304 phiên và vàng futures GC=F có 1.257 phiên, từ 15/10/2021 đến 14/10/2026. GC=F không được đổi tên thành spot XAU/USD. Lợi suất được tính theo phiên có thật, không nội suy ngày nghỉ. Dữ liệu e-commerce không được đưa vào thực nghiệm mới vì chưa xác minh được thứ tự sự kiện và nhãn đủ để biện minh cho RNN.',RED)
para('Ngày của nhãn quyết định tập dữ liệu: FPT train trước 13/02/2026, validation từ 13/02 đến trước 04/06/2026; vàng train trước 26/02/2026, validation đến trước 09/06/2026. Test kéo dài đến nhãn 14/10/2026. Cách định nghĩa này khác chỉ số ngày đặc trưng của bảng thị trường ở Chương 5; không so sánh trực tiếp hai bảng như cùng tập test. Validation/test được phép dùng lịch sử đã biết trước ngày nhãn, phù hợp dự báo cuốn chiếu một bước.',RED)
table(['Miền','Train / Val / Test','Khoảng ngày nhãn test'],[[r['domain'],f"{r['train_size']} / {r['validation_size']} / {r['test_size']}",f"{r['test_start']} — {r['test_end']}"] for r in SEQ['rows'] if r['model']=='RNN'], 'Phân chia chuỗi theo ngày nhãn',RED,[3,5,8])
para('Ba mạng dùng 16 hidden units, một recurrent layer và một linear head, Adam learning_rate=0,001, 35 epoch, batch_size=128 và seed=42. Trung bình và độ lệch chuẩn của đầu vào/nhãn chỉ tính trên train. Checkpoint tối thiểu validation MSE được lưu; mạng triển khai được chọn bằng validation MAE theo đơn vị giá, có Persistence cùng tham gia. Test được giữ cho báo cáo cuối.',RED)
para('GRU trong PyTorch có hai khác biệt quy ước so với công thức lý thuyết ở trên: reset gate nhân sau phép biến đổi tuyến tính của hidden state trong nhánh ứng viên; update gate z đặt trọng số lên trạng thái cũ thay vì ứng viên. Cụ thể, hₜ=(1−zₜ)⊙nₜ+zₜ⊙hₜ₋₁, với nₜ là ứng viên. API dùng đúng hai quy ước của checkpoint; kiểm thử cả sáu mạng phát hiện sai khác thứ tự cổng, vị trí reset và cách phối hợp trạng thái.',RED)
add_details('sequence_mechanism',globals())
heading('4.3. So sánh RNN, LSTM, GRU và Persistence',2,RED)
table(['Miền / Mạng','Tham số','Val MAE','Test MAE','Test RMSE'],[[r['domain']+' / '+r['model'],str(r['parameters']),f"{r['validation_MAE']:.3f}",f"{r['MAE']:.3f}",f"{r['RMSE']:.3f}"] for r in SEQ['rows']], 'Kết quả mạng chuỗi trong đơn vị gốc VND hoặc USD/oz',RED,[4,2,3.3,3.3,3.4])
for domain,v in SEQ['deployment'].items():
    para(f"{domain.upper()}: Persistence test MAE {v['baseline_test_metrics']['MAE']:.3f} {v['unit']}. Mạng có validation MAE thấp nhất là {v['neural_validation_winner']} ({v['neural_validation_MAE']:.3f}), trong khi Persistence validation MAE {v['baseline_validation_MAE']:.3f}. Vì vậy quy tắc trước test chọn {v['kind']} để trình diễn trên web.",RED)
picture(FIG/'fpt_sequences.png','Dự báo một bước FPT bằng ba mạng và baseline',RED,max_height=7.5)
picture(FIG/'gold_sequences.png','Dự báo một bước vàng futures GC=F',RED,max_height=7.5)
para('Vanilla RNN có MAE test thấp nhất trong ba mạng ở cả hai miền trong lần chạy này, nhưng không được đổi lựa chọn triển khai sau khi nhìn test. Kết quả bác bỏ việc mặc định LSTM luôn tốt hơn. Chuỗi ngắn, ít dữ liệu, nhiễu thị trường và một random seed hạn chế khả năng khái quát hóa. Mức chênh MAE nhỏ chưa chứng minh lợi ích giao dịch; báo cáo không suy ra lợi nhuận.',RED)
add_details('sequence_errors',globals())
heading('4.4. Hạn chế và khả năng mở rộng',2,RED)
para('Thực nghiệm là dự báo một bước trên snapshot cố định, không phải dự báo nhiều bước tương lai hoặc dữ liệu trực tiếp. Phân tích bootstrap theo khối ở trên chưa thay thế kiểm tra nhiều seed và walk-forward backtest. Hướng mở rộng gồm nghiên cứu độ nhạy với chiều dài cửa sổ, các biến đầu vào có timestamp và ngân sách tuning khóa trước. Không dùng test lặp lại để điều chỉnh cấu hình.',RED)
para('Mô hình triển khai giữ checkpoint train lịch sử, không refit bằng test để giữ được dấu vết đánh giá. Người dùng có thể sửa giá gần nhất và 30 log-return để thử kịch bản; kết quả không có khoảng tin cậy đã hiệu chỉnh cho kịch bản ngoài miền. Checkpoint của mọi RNN/LSTM/GRU và training history được lưu riêng để kiểm tra.',RED)

start_chapter(5,'KNOWLEDGE GRAPH VÀ RAG')
add_intro(5,globals())
# Lấy nguyên các đoạn văn của TL02, loại header/footer, citation và tham chiếu cũ.
ch5_figs={4:('architecture.png','Kiến trúc KG và RAG'),8:('data_profile.png','Thông tin ba miền dữ liệu'),9:('knowledge_graph.png','Đồ thị tri thức đa miền'),11:('code_kg_core.png','Cấu trúc KnowledgeGraphCore'),14:('architecture.png','Các bước truy xuất và tổng hợp bằng chứng'),15:('code_retriever.png','Lập chỉ mục và truy xuất TF-IDF'),16:('code_graphrag.png','Truy xuất bằng chứng đồ thị'),18:('rag_benchmark.png','Đánh giá smoke test RAG và GraphRAG'),19:('diabetes_model.png','Mô hình hồi quy tiến triển tiểu đường'),20:('kg_diabetes.png','Đồ thị tri thức y tế'),21:('code_diabetes_training.png','Huấn luyện và đánh giá mô hình y tế'),23:('fpt_market.png','Chuỗi giá và chỉ báo FPT'),24:('code_market_data.png','Pipeline xử lý dữ liệu thị trường'),25:('stock_model.png','Dự báo FPT so với Persistence'),26:('kg_stock.png','Đồ thị tri thức chứng khoán'),28:('gold_market.png','Chuỗi giá và chỉ báo vàng'),29:('kg_gold.png','Đồ thị tri thức vàng'),30:('gold_model.png','Mô hình vàng và baseline'),31:('code_market_training.png','Quy trình học máy cho thị trường'),33:('web_app_preview.png','Sơ đồ bố cục giao diện gốc'),34:('code_deployment.png','Xuất estimator sang JavaScript')}
for page in S2[3:37]:
    pg=page['page'];parts=[]
    for b in page['blocks']:
        if b['type']!='text':continue
        text=clean(b['text'])
        if b['bbox'][1]<55 or re.fullmatch(r'\d+',text) or text.startswith('Hình:'):continue
        if text.startswith('CHƯƠNG 4 -') or text.startswith('Kiều Trường Giang -'):continue
        if text in ['ĐÁNH GIÁ TỔNG HỢP','KẾT LUẬN']:
            if text=='KẾT LUẬN':continue
            heading('5.8. Đánh giá tổng hợp',2);continue
        hm=re.match(r'^(4\.\d+(?:\.\d+)?)\.\s+(.+)',text)
        if hm:
            num,title=hm.groups();num='5'+num[1:]
            if num=='5.1.1' and 'RAG' in title:num='5.1.2'
            elif num=='5.1.2':num='5.1.3'
            elif num=='5.1.3':num='5.1.4'
            heading(num+'. '+title,3 if num.count('.')==2 else 2);continue
        if pg==37:continue  # Kết luận sẽ đặt sau các bảng tổng hợp.
        color=BLACK
        if text.startswith('Web:'):
            link_line('Web','https://rag-httm-chatbot.vercel.app/',RED)
            link_line('GitHub','https://github.com/jellalaz/HTTM_ThayQue',RED)
            para('Hình bố cục dưới đây là sơ đồ minh họa, không phải ảnh chụp màn hình.',RED)
            continue
        if 'GitHub: https://github.com/jellalaz/RAG_HTTM' in text:
            text=text.replace('https://github.com/jellalaz/RAG_HTTM','https://github.com/jellalaz/HTTM_ThayQue');color=RED
        # Khẳng định nguồn tải trực tiếp được cập nhật cho lần chạy snapshot hiện tại.
        if 'được tải trực tiếp khi chạy' in text:
            text=text.replace('được tải trực tiếp khi chạy','được đọc từ snapshot đã khóa trong lần tái lập; có chế độ tải mới khi được bật');color=RED
        para(text,color,f'TL02 trang {pg}')
    if pg in ch5_figs:
        file,cap=ch5_figs[pg]
        if pg==14:picture(FIG/'mechanism_rag.png','Minh họa luồng truy xuất và tổng hợp bằng chứng',RED,max_height=5)
        else:picture(ROOT/'artifacts_giangkt/figures'/file,cap,BLACK,max_height=7.3,source=f'TL02 trang {pg}')

heading('5.8.1. Kết quả huấn luyện ba miền',3)
legacy=json.loads((ROOT/'artifacts_giangkt/data/web_model_config.json').read_text())
table(['Miền','Mô hình','CV MAE','Test MAE','R²'],[[key,m['name'],f"{m['metrics']['CV_MAE']:.3f}",f"{m['metrics']['MAE']:.3f}",f"{m['metrics']['R2']:.3f}"] for key,m in legacy['models'].items()], 'Tổng hợp kết quả giữ từ Tiểu luận 02',widths=[3,4.5,2.8,2.8,2.9])
para('Mô hình được chọn theo CV MAE trên tập train. Chỉ số test được tính trên phần dữ liệu giữ lại.',BLACK,'TL02 trang 38')
heading('5.8.2. Kết quả truy xuất và bộ ba tri thức',3)
rows=list(csv.DictReader((ROOT/'artifacts_giangkt/data/retrieval_evaluation.csv').open()))
table(['Câu hỏi','Miền','P@3','Graph','Hybrid'],[[r['query'],r['expected_domain'],f"{float(r['vector_precision@3']):.2f}",r['graph_evidence'],f"{float(r['hybrid_evidence_score']):.2f}"] for r in rows], 'Kết quả smoke test được giữ nguyên',widths=[8,2.4,1.8,1.8,2])
para('Bảng này là phép kiểm tra hồi quy cho pipeline. Nó không thay thế đánh giá câu trả lời bởi chuyên gia.',BLACK,'TL02 trang 39')
rows=list(csv.DictReader((ROOT/'artifacts_giangkt/data/knowledge_graph_triples.csv').open()))
table(['Subject','Predicate','Object'],[[r['Subject'],r['Predicate'],r['Object']] for r in rows], 'Bộ ba tri thức của hệ thống',widths=[5.2,5.8,5])
heading('5.8.3. Chunking, embedding và giới hạn đánh giá',3,RED)
para('Embedding là biểu diễn văn bản bằng vector; dense embedding học ngữ nghĩa từ dữ liệu huấn luyện, còn TF-IDF là vector thưa theo tần suất từ. Vector database bổ sung chỉ mục và quản trị truy vấn ở quy mô lớn; prototype hiện lưu ma trận TF-IDF trong bộ nhớ, chưa vận hành một vector database độc lập. Corpus gồm các tài liệu ngắn đã là đơn vị truy xuất, nên không tạo hàng loạt chunk chồng lấn không cần thiết. Với PDF dài, cần chunk theo heading, giữ metadata và thử kích thước/overlap bằng một tập câu hỏi riêng.',RED)
para('Điểm hybrid cũ cộng 0,25 khi có bằng chứng đồ thị là điểm kỹ thuật tự định nghĩa, không phải Accuracy, faithfulness hay xác suất đúng của câu trả lời. Precision theo domain chỉ đo mức đúng miền. Việc mọi câu hỏi tìm được một cạnh không chứng minh cạnh đó đủ để trả lời. Cần phân biệt truy xuất văn bản, độ đúng đường đi và độ bám sát bằng chứng khi đánh giá generator.',RED)
heading('5.8.4. Đồng bộ notebook và các chức năng web bổ sung',3,RED)
para('main.ipynb chứa các bước thực nghiệm mới, pipeline KG/RAG từ Tiểu luận 02, xuất mô hình và đóng gói web. Chế độ full là cấu hình của số liệu trong báo cáo; quick dùng ít dữ liệu/epoch hơn để kiểm tra, không thay số liệu full trong báo cáo. Manifest lưu seed, phiên bản thư viện, hash dataset, checkpoint và cấu hình chạy.',RED)
para('Giao diện Insight Lab giữ khu dự đoán ba miền, trợ lý sức khỏe, màu sắc và bố cục chính. Khu thực nghiệm bổ sung giá nhà, phân loại nhãn nguy cơ tổng hợp, nhận dạng chữ số CNN và dự báo chuỗi. Các mô hình mới sử dụng API /api/predict và trọng số JSON xuất từ notebook. API không gọi dịch vụ LLM, không cần secret và có kiểm tra kích thước cũng như giá trị đầu vào. Chatbot sức khỏe cũ vẫn xử lý trong trình duyệt; dữ liệu nhập vào các demo API được gửi tới máy chủ inference để tính dự đoán.',RED)
add_details('rag_trace',globals())
heading('5.9. Kết luận chương',2)
for b in S2[36]['blocks']:
    if b['type']=='text' and 100<b['bbox'][1]<730 and clean(b['text'])!='KẾT LUẬN' and not re.fullmatch(r'\d+',clean(b['text'])):para(b['text'],BLACK,'TL02 trang 37')

# Page number fields and first-page suppression.
section.different_first_page_header_footer=True
p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.save(ROOT/f'{NAME}.docx')
(DATA/'report_provenance.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2),encoding='utf-8')
(DATA/'report_structure.json').write_text(json.dumps({'headings':headings,'figures':dict(fig_count),'tables':dict(table_count),'equations':dict(eq_count)},ensure_ascii=False,indent=2),encoding='utf-8')
print('Created',NAME+'.docx',len(doc.paragraphs),'paragraphs',len(doc.tables),'tables')
